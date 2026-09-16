from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modules.official_files import CUMULATIVE_NAME, PENDING_NAME, discover_official_files
from self_healing_main import _build_baseline_files


class OfficialFileDiscoveryTests(unittest.TestCase):
    def test_discovers_exact_official_names(self):
        items = [
            {"id": "c1", "name": CUMULATIVE_NAME, "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            {"id": "p1", "name": PENDING_NAME, "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            {"id": "x1", "name": "累计加工台账-备份.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        ]
        result = discover_official_files(items)
        self.assertEqual(result[CUMULATIVE_NAME]["id"], "c1")
        self.assertEqual(result[PENDING_NAME]["id"], "p1")

    def test_missing_official_file_is_none(self):
        result = discover_official_files([])
        self.assertIsNone(result[CUMULATIVE_NAME])
        self.assertIsNone(result[PENDING_NAME])

    def test_duplicate_official_name_blocks(self):
        items = [
            {"id": "c1", "name": CUMULATIVE_NAME, "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            {"id": "c2", "name": CUMULATIVE_NAME, "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        ]
        with self.assertRaisesRegex(RuntimeError, "多个同名正式文件"):
            discover_official_files(items)


class BaselineWorkbookTests(unittest.TestCase):
    def test_baseline_uses_current_generator_and_starts_from_zero(self):
        source_records = [{
            "order_raw": "THP10-8000J-0911",
            "order": "THP10-8000J-0911",
            "drawing": "2310",
            "thickness": 80.0,
            "quantity": 4,
            "length": 1000.0,
            "width": 500.0,
            "bevel": "W",
            "total_weight_t": 1.256,
            "source_file": "汇总表.xlsx",
            "order_container_name": "THP10-8000J-0911",
            "source_row": 2,
        }]
        with tempfile.TemporaryDirectory() as temp_dir:
            cumulative_path, pending_path = _build_baseline_files(source_records, Path(temp_dir))
            cumulative = load_workbook(cumulative_path, read_only=True, data_only=True)
            self.assertEqual(
                cumulative.sheetnames,
                ["累计加工台账", "加工流水", "板材入账记录", "异常记录"],
            )
            pending = load_workbook(pending_path, read_only=True, data_only=True)
            self.assertEqual(pending.sheetnames, ["当前待加工零件"])

            ws = cumulative["累计加工台账"]
            values = [cell.value for row in ws.iter_rows() for cell in row]
            self.assertIn("2310", values)
            # 明细列 I/J 分别为“累计已加工/当前剩余”。
            detail_rows = [row for row in ws.iter_rows(min_row=1, values_only=True) if row[0] == "2310"]
            self.assertEqual(len(detail_rows), 1)
            self.assertEqual(detail_rows[0][8], 0)
            self.assertEqual(detail_rows[0][9], 4)


if __name__ == "__main__":
    unittest.main()
