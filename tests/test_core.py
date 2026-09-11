import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook

from modules.excel_generator import generate_cumulative_report, generate_pending_report
from modules.excel_reader import (
    normalize_order_key,
    read_existing_ledger,
    read_source_summary,
    read_split_result,
)
from modules.idempotency import reconcile_posted_board
from modules.archive_manager import should_archive
from modules.process_parts import build_current_state, validate_new_board
from modules.process_parts import board_content_fingerprint


class CoreWorkflowTests(unittest.TestCase):
    def _make_source(self, path: Path, *, kg=False, blank_bevel_header=False):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "订单汇总表"
        headers = [
            "订单号", "图号", "厚度", "件数", "长(mm)", "宽(mm)",
            "切割长度(mm)", "净面积(m²)", "" if blank_bevel_header else "坡口",
            "总净重(kg)" if kg else "总净重(t)", "核对备注",
        ]
        for c, value in enumerate(headers, 1):
            ws.cell(4, c, value)
        ws.append(["175.26-08-05  D53K-1600A-0805", "D53K-1600A.1-1-13", 150, 2, 1980, 325, 0, 0, "W", 1515.4425 if kg else 1.5154425, ""])
        wb.save(path)

    def _make_split(self, path: Path, *, order="D53K-1600A-0805", drawing="1-1-13"):
        wb = Workbook()
        ws = wb.active
        headers = ["序号", "订单号", "图号", "厚度", "基础件数", "图片拆分数量", "基础表总重量（T）", "单件重量（T）", "本次拆分重量（KG）"]
        for c, value in enumerate(headers, 1):
            ws.cell(3, c, value)
        ws.append([1, order, drawing, 150, 2, 2, 1.5154425, 0.75772125, 1515.4425])
        ws["K7"] = "图片标注重量（KG）"
        ws["M7"] = 1516.0
        ws["K8"] = "拆分重量合计（KG）"
        ws["M8"] = 1515.4425
        ws["K9"] = "差额绝对值（KG）"
        ws["M9"] = 0.5575
        ws["K10"] = "复核结论"
        ws["M10"] = "通过"
        wb.save(path)

    def _make_ledger(self, path: Path, *, completed=False, board_id=""):
        wb = Workbook()
        ws = wb.active
        ws.title = "累计加工台账"
        ws["A1"] = "正在进行加工订单"
        remaining = 0 if completed else 2
        processed = 2 if completed else 0
        pending_weight = 0 if completed else 1.5154425
        status = "已完成" if completed else "未开始"
        ws["A6"] = f"订单：D53K-1600A-0805 ｜ {'已完成' if completed else '加工中'} ｜ 待加工 {remaining} 件 ｜ 待加工 {pending_weight:.3f} t"
        headers = ["图号", "厚度(mm)", "坡口", "长(mm)", "宽(mm)", "订单总数量", "零件总重量(t)", "板材号/加工来源", "累计已加工", "当前剩余", "待加工重量(t)", "零件状态"]
        for c, value in enumerate(headers, 1):
            ws.cell(7, c, value)
        ws.append(["D53K-1600A.1-1-13", 150, "W", 1980, 325, 2, 1.5154425, f"{board_id}×2" if board_id else "", processed, remaining, pending_weight, status])

        f = wb.create_sheet("加工流水")
        f.append(["板材号", "拆图结果文件", "图号", "厚度(mm)", "坡口", "本张板加工数量", "数据性质", "说明"])
        b = wb.create_sheet("板材入账记录")
        b.append(["数据性质", "板材号", "拆图结果文件", "计入件数", "状态", "归档位置"])
        a = wb.create_sheet("异常记录")
        a.append(["板材号", "业务日期", "订单号", "图号", "厚度(mm)", "数量", "异常类型", "说明"])
        if board_id:
            f.append([board_id, f"{board_id}_完成.xlsx", "D53K-1600A.1-1-13", 150, "W", 2, "正式入账", "测试"])
            b.append(["正式入账", board_id, f"{board_id}_完成.xlsx", 2, "已核验并入账", "拆图结果/已录入数量"])
        wb.save(path)

    def test_source_parser_kg_to_tons(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.xlsx"
            self._make_source(path, kg=True)
            rows = read_source_summary(path)
            self.assertEqual(rows[0]["order"], "D53K-1600A-0805")
            self.assertAlmostEqual(rows[0]["total_weight_t"], 1.5154425)

    def test_source_parser_infers_unique_unlabelled_bevel_column(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "legacy-template.xlsm"
            self._make_source(path, blank_bevel_header=True)
            rows = read_source_summary(path)
            self.assertEqual(rows[0]["bevel"], "W")

    def test_source_parser_accepts_status_suffix_in_order_folder(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "YT71S-2500Z-0715 模板.xlsm"
            self._make_source(path)
            wb = load_workbook(path)
            ws = wb.active
            ws["A5"] = "YT71S-2500Z-0715"
            wb.save(path)

            rows = read_source_summary(
                path,
                source_name=path.name,
                order_container_name="159.26-07-15  YT71S-2500Z-0715 已做完核算表",
            )
            self.assertEqual(rows[0]["order"], "YT71S-2500Z-0715")

    def test_order_folder_uses_unique_order_number_not_status_text(self):
        self.assertEqual(
            normalize_order_key("182.26-08-25  THP11-10000Q-0825 已做完核算表"),
            "THP11-10000Q-0825",
        )

    def test_order_folder_ignores_spaces_around_hyphen(self):
        self.assertEqual(
            normalize_order_key("190.26-09-11   THP10-8000J -0911 已做完核算表 待审"),
            "THP10-8000J-0911",
        )

    def test_same_board_number_uses_content_to_distinguish_materials(self):
        first = {"rows": [{
            "order": "YT71S-2500Z-0715", "drawing": "A1", "thickness": 50,
            "bevel": "", "base_quantity": 10, "split_quantity": 2,
            "base_total_weight_t": 1.25,
        }]}
        same = deepcopy(first)
        different = deepcopy(first)
        different["rows"][0]["split_quantity"] = 3
        self.assertEqual(board_content_fingerprint(first), board_content_fingerprint(same))
        self.assertNotEqual(board_content_fingerprint(first), board_content_fingerprint(different))

    def test_existing_ledger_exposes_permanent_historical_rows(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "ledger.xlsx"
            self._make_ledger(path, completed=True, board_id="#309")
            existing = read_existing_ledger(path)
            self.assertEqual(len(existing["historical_rows"]), 1)
            self.assertEqual(existing["historical_rows"][0]["remaining"], 0)
            self.assertEqual(existing["historical_rows"][0]["board_sources"], "#309×2")
            self.assertIn("#309", existing["posted_boards"])

    def test_archive_requires_post_write_verification(self):
        self.assertTrue(should_archive(True))
        self.assertFalse(should_archive(False))
        self.assertFalse(should_archive("完成"))

    def test_split_and_d53k_unique_suffix_match(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            split_path = Path(td) / "#309_完成.xlsx"
            ledger_path = Path(td) / "ledger.xlsx"
            self._make_source(source_path)
            self._make_split(split_path)
            self._make_ledger(ledger_path)
            sources = read_source_summary(source_path)
            payload = read_split_result(split_path, board_id="#309")
            existing = read_existing_ledger(ledger_path)
            result = validate_new_board(
                board_id="#309",
                filename="#309_完成.xlsx",
                split_payload=payload,
                source_records=sources,
                posted_boards=existing["posted_boards"],
            )
            self.assertTrue(result["ok"], result.get("error"))
            self.assertEqual(result["board_record"]["计入件数"], 2)
            self.assertIn("D53K", result["flows"][0]["图号"])

    def test_unmatched_order_blocks_entire_board(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            split_path = Path(td) / "2319_完成.xlsx"
            self._make_source(source_path)
            self._make_split(split_path, order="D53K-8000D-0718", drawing="05TD1")
            sources = read_source_summary(source_path)
            payload = read_split_result(split_path, board_id="2319")
            result = validate_new_board(
                board_id="2319",
                filename="2319_完成.xlsx",
                split_payload=payload,
                source_records=sources,
                posted_boards=set(),
            )
            self.assertFalse(result["ok"])
            self.assertEqual(result["anomaly"]["异常类型"], "阻断入账")

    def test_posted_board_reconciliation_prevents_double_count(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            split_path = Path(td) / "#309_完成.xlsx"
            ledger_path = Path(td) / "ledger.xlsx"
            self._make_source(source_path)
            self._make_split(split_path)
            self._make_ledger(ledger_path, completed=True, board_id="#309")
            sources = read_source_summary(source_path)
            payload = read_split_result(split_path, board_id="#309")
            existing = read_existing_ledger(ledger_path)
            recovery = reconcile_posted_board(
                board_id="#309",
                filename="#309_完成.xlsx",
                split_payload=payload,
                source_records=sources,
                existing_flows=existing["flows"],
                existing_board_records=existing["board_records"],
            )
            self.assertTrue(recovery["ok"], recovery.get("reason"))
            self.assertEqual(recovery["expected_qty"], 2)

    def test_state_and_reports(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            self._make_source(source_path)
            sources = read_source_summary(source_path)
            key = ("D53K-1600A-0805", "D53K-1600A.1-1-13", 150.0, "W")
            state = build_current_state(
                sources,
                existing_state={key: {"processed": 0, "board_sources": ""}},
                delta_by_key={key: 1},
                source_additions={key: [("#309", 1)]},
            )
            self.assertEqual(state[0]["remaining"], 1)
            self.assertEqual(state[0]["status"], "部分完成")
            self.assertEqual(state[0]["board_sources"], "#309×1")

            cumulative = Path(td) / "累计加工台账.xlsx"
            pending = Path(td) / "当前待加工零件.xlsx"
            generate_cumulative_report(state, [], [], [], cumulative, note="测试")
            generate_pending_report(state, pending, note="测试")
            cwb = load_workbook(cumulative)
            self.assertEqual(set(cwb.sheetnames), {"累计加工台账", "加工流水", "板材入账记录", "异常记录"})
            pws = load_workbook(pending)["当前待加工零件"]
            j_cells = [c for c in pws["J"] if c.value == 1]
            self.assertTrue(j_cells)
            self.assertTrue(j_cells[0].font.bold)
            self.assertEqual(j_cells[0].fill.fgColor.rgb[-6:], "FFF2CC")

    def test_status_fill_covers_entire_detail_row(self):
        rows = []
        statuses = [
            ("已完成", 2, 0, 0.0, "E2F0D9"),
            ("部分完成", 1, 1, 0.5, "FFF2CC"),
            ("超加工/待核查", 3, -1, -0.5, "F4CCCC"),
        ]
        for index, (status, processed, remaining, pending_weight, _) in enumerate(statuses, start=1):
            rows.append({
                "order": "STATUS-COLOR-TEST",
                "drawing": f"COLOR-{index}",
                "thickness": 60,
                "bevel": "",
                "length": 1000,
                "width": 500,
                "quantity": 2,
                "total_weight_t": 1.0,
                "board_sources": "TEST",
                "processed": processed,
                "remaining": remaining,
                "pending_weight_t": pending_weight,
                "status": status,
            })

        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "状态颜色整行测试.xlsx"
            generate_cumulative_report(rows, [], [], [], output, note="状态颜色回归测试")
            ws = load_workbook(output)["累计加工台账"]

            detail_rows = {
                ws.cell(row, 1).value: row
                for row in range(1, ws.max_row + 1)
                if isinstance(ws.cell(row, 1).value, str) and ws.cell(row, 1).value.startswith("COLOR-")
            }
            self.assertEqual(set(detail_rows), {"COLOR-1", "COLOR-2", "COLOR-3"})

            for index, (_, _, _, _, expected_color) in enumerate(statuses, start=1):
                row = detail_rows[f"COLOR-{index}"]
                actual_colors = [ws.cell(row, col).fill.fgColor.rgb[-6:] for col in range(1, 13)]
                self.assertEqual(actual_colors, [expected_color] * 12)
                self.assertTrue(ws.cell(row, 10).font.bold)
                self.assertEqual(ws.cell(row, 10).font.sz, 14)


if __name__ == "__main__":
    unittest.main()
