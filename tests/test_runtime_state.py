import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modules.excel_generator import RED, YELLOW, generate_pending_report
from modules.process_parts import build_current_state, order_is_complete, board_content_fingerprint
from modules.runtime_state import RuntimeState, STATE_SCHEMA_VERSION, deserialize_ledger, serialize_ledger, source_signature


def source(order="O-1", drawing="D-1", qty=3, thickness=20.0, bevel=""):
    return {"order_raw": order, "order": order, "drawing": drawing, "thickness": thickness, "quantity": qty,
            "length": 100.0, "width": 50.0, "bevel": bevel, "total_weight_t": 0.3,
            "source_file": "汇总表.xlsx", "order_container_name": order, "source_row": 2}


def item(fid="f1", modified="2026-01-01T00:00:00Z", md5="abc", parent="p1", container="O-1"):
    return {"id": fid, "name": "汇总表.xlsx", "modifiedTime": modified, "md5Checksum": md5, "size": "100",
            "parent_folder_id": parent, "order_folder_id": parent, "order_container_name": container}


class RuntimeStateTests(unittest.TestCase):
    def test_missing_cache_full_initialization_marks_all_added(self):
        with tempfile.TemporaryDirectory() as td:
            state = RuntimeState(Path(td) / "state.json")
            c = state.classify_sources([item(), item("f2")])
            self.assertEqual(len(c["added"]), 2); self.assertFalse(state.valid)

    def test_unchanged_uses_full_metadata_signature(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "state.json"; state = RuntimeState(p); state.set_source_records(item(), [source()]); state.save()
            restored = RuntimeState(p); c = restored.classify_sources([item()])
            self.assertEqual(len(c["unchanged"]), 1); self.assertEqual(restored.get_source(item())["records"][0]["drawing"], "D-1")

    def test_md5_parent_or_container_change_is_modified(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "state.json"; state = RuntimeState(p); state.set_source_records(item(), [source()]); state.save()
            for changed in [item(md5="def"), item(parent="p2"), item(container="O-2")]:
                c = RuntimeState(p).classify_sources([changed]); self.assertEqual(len(c["modified"]), 1)

    def test_removed_source_is_detected_and_pruned(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"state.json"; state=RuntimeState(p); state.set_source_records(item(),[source()]); state.save()
            restored=RuntimeState(p); self.assertEqual(len(restored.classify_sources([])["removed"]),1); restored.retain_sources([]); self.assertEqual(restored.sources,{})

    def test_schema_mismatch_invalidates_cache(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"state.json"; p.write_text(json.dumps({"state_schema_version": STATE_SCHEMA_VERSION+1,"sources":{}}),encoding="utf-8")
            self.assertFalse(RuntimeState(p).valid)

    def test_ledger_round_trip_preserves_tuple_keys_and_sets(self):
        ledger={"state":{("O","D",20.0,""):{"processed":2,"board_sources":"#1×2"}},"historical_rows":[],"flows":[],"board_records":[],"anomalies":[],"posted_boards":{"#1"},"posted_board_keys":{("#1","fp")},"legacy_posted_boards":set()}
        restored=deserialize_ledger(serialize_ledger(ledger)); self.assertEqual(restored["state"][("O","D",20.0,"")]["processed"],2); self.assertIn(("#1","fp"),restored["posted_board_keys"])

    def test_ledger_cache_requires_same_file_and_md5(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"state.json"; state=RuntimeState(p); state.set_ledger({"id":"ledger","modifiedTime":"x","md5Checksum":"m1","size":"1"},{"state":{},"historical_rows":[],"flows":[],"board_records":[],"anomalies":[],"posted_boards":set(),"posted_board_keys":set(),"legacy_posted_boards":set()}); state.save()
            restored=RuntimeState(p); self.assertTrue(restored.ledger_matches({"id":"ledger","md5Checksum":"m1"})); self.assertFalse(restored.ledger_matches({"id":"ledger","md5Checksum":"m2"}))


class BusinessBoundaryTests(unittest.TestCase):
    def test_overprocess_keeps_negative_and_red_status(self):
        rows=build_current_state([source(qty=3)],{("O-1","D-1",20.0,""):{"processed":2,"board_sources":"#old×2"}},{("O-1","D-1",20.0,""):2},{("O-1","D-1",20.0,""):[("#new",2)]})
        self.assertEqual(rows[0]["processed"],4); self.assertEqual(rows[0]["remaining"],-1); self.assertEqual(rows[0]["status"],"超加工/待核查")
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"pending.xlsx"; generate_pending_report(rows,p); wb=load_workbook(p); ws=wb["当前待加工零件"]
            status_row=next(r for r in ws.iter_rows() if r[11].value=="超加工/待核查")
            self.assertEqual(status_row[0].fill.fgColor.rgb[-6:],RED)

    def test_unprocessed_status_name_and_yellow(self):
        rows=build_current_state([source()],{}); self.assertEqual(rows[0]["status"],"未加工")
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"pending.xlsx"; generate_pending_report(rows,p); ws=load_workbook(p)["当前待加工零件"]
            status_row=next(r for r in ws.iter_rows() if r[11].value=="未加工"); self.assertEqual(status_row[0].fill.fgColor.rgb[-6:],YELLOW)

    def test_order_exits_only_when_all_rows_zero(self):
        a=build_current_state([source(drawing="A",qty=1),source(drawing="B",qty=1)],{("O-1","A",20.0,""):{"processed":1},("O-1","B",20.0,""):{"processed":0}})
        self.assertFalse(order_is_complete(a)); b=build_current_state([source(drawing="A",qty=1),source(drawing="B",qty=1)],{("O-1","A",20.0,""):{"processed":1},("O-1","B",20.0,""):{"processed":1}}); self.assertTrue(order_is_complete(b))

    def test_fingerprint_same_content_ignores_row_order(self):
        r1={"order":"O","drawing":"A","thickness":20,"bevel":"","base_quantity":3,"split_quantity":1,"base_total_weight_t":1.2}; r2={**r1,"drawing":"B"}
        self.assertEqual(board_content_fingerprint({"rows":[r1,r2]}),board_content_fingerprint({"rows":[r2,r1]}))


if __name__ == "__main__": unittest.main()
