from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modules.excel_generator import generate_cumulative_report
from modules.process_parts import build_current_state, validate_new_board


class BusinessRulesRegressionTests(unittest.TestCase):
    def _source(self):
        return [{
            "order": "ORDER-1",
            "drawing": "A1",
            "thickness": 20.0,
            "bevel": "",
            "length": 1000.0,
            "width": 500.0,
            "quantity": 3,
            "total_weight_t": 0.3,
        }]

    def _split(self, qty=1):
        return {"rows": [{
            "order": "ORDER-1",
            "drawing": "A1",
            "thickness": 20.0,
            "bevel": "",
            "base_quantity": 3,
            "split_quantity": qty,
            "base_total_weight_t": 0.3,
        }]}

    def _detail_row(self, output: Path) -> tuple:
        ws = load_workbook(output)["累计加工台账"]
        row = next(
            r for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value == "A1"
        )
        return ws, row

    def test_same_full_board_id_with_different_content_is_blocked(self):
        result = validate_new_board(
            board_id="#2326",
            filename="#2326_完成.xlsx",
            split_payload=self._split(qty=2),
            source_records=self._source(),
            posted_boards={"#2326"},
        )
        self.assertFalse(result["ok"])
        self.assertIn("同板材号内容冲突", result["error"])
        self.assertEqual(result["anomaly"]["异常类型"], "同板材号内容冲突")

    def test_new_board_is_still_allowed(self):
        result = validate_new_board(
            board_id="#2326-1",
            filename="#2326-1_完成.xlsx",
            split_payload=self._split(qty=1),
            source_records=self._source(),
            posted_boards={"#2326"},
        )
        self.assertTrue(result["ok"], result.get("error"))

    def test_unprocessed_status_has_no_fill(self):
        state = build_current_state(self._source(), existing_state={})
        self.assertEqual(state[0]["status"], "未加工")
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "status.xlsx"
            generate_cumulative_report(state, [], [], [], output)
            ws, row = self._detail_row(output)
            self.assertTrue(all(ws.cell(row, col).fill.fill_type is None for col in range(1, 13)))

    def test_partially_processed_status_is_yellow(self):
        key = ("ORDER-1", "A1", 20.0, "")
        state = build_current_state(
            self._source(),
            existing_state={key: {"processed": 1, "board_sources": "#100×1"}},
        )
        self.assertEqual(state[0]["status"], "部分完成")
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "status.xlsx"
            generate_cumulative_report(state, [], [], [], output)
            ws, row = self._detail_row(output)
            colors = [ws.cell(row, col).fill.fgColor.rgb[-6:] for col in range(1, 13)]
            self.assertEqual(colors, ["FFF2CC"] * 12)

    def test_overprocessing_keeps_negative_remaining(self):
        key = ("ORDER-1", "A1", 20.0, "")
        state = build_current_state(
            self._source(),
            existing_state={key: {"processed": 2, "board_sources": "#100×2"}},
            delta_by_key={key: 2},
        )
        self.assertEqual(state[0]["processed"], 4)
        self.assertEqual(state[0]["remaining"], -1)
        self.assertEqual(state[0]["status"], "超加工/待核查")


if __name__ == "__main__":
    unittest.main()
