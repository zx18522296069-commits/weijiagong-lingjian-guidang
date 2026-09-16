from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modules.excel_generator import generate_cumulative_report


class OverprocessingAnomalyTests(unittest.TestCase):
    @staticmethod
    def _negative_row():
        return {
            "order": "ORDER-OVER",
            "drawing": "A1",
            "thickness": 20,
            "bevel": "W",
            "length": 1000,
            "width": 500,
            "quantity": 3,
            "total_weight_t": 1.5,
            "board_sources": "#1×2；#2×2",
            "processed": 4,
            "remaining": -1,
            "pending_weight_t": 0,
            "status": "超加工/待核查",
        }

    def test_negative_remaining_is_red_and_written_to_anomaly_sheet(self):
        row = self._negative_row()
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "累计加工台账.xlsx"
            generate_cumulative_report([row], [], [], [], path)
            wb = load_workbook(path)
            ws = wb["累计加工台账"]
            detail = next(r for r in ws.iter_rows() if r[0].value == "A1")
            self.assertEqual(detail[9].value, -1)
            self.assertEqual(detail[11].value, "超加工/待核查")
            self.assertEqual([cell.fill.fgColor.rgb[-6:] for cell in detail[:12]], ["F4CCCC"] * 12)

            anomaly_ws = wb["异常记录"]
            headers = [cell.value for cell in anomaly_ws[1]]
            records = [dict(zip(headers, [cell.value for cell in r])) for r in anomaly_ws.iter_rows(min_row=2)]
            over = [r for r in records if r.get("异常类型") == "超加工/待核查"]
            self.assertEqual(len(over), 1)
            self.assertEqual(over[0]["订单号"], "ORDER-OVER")
            self.assertEqual(over[0]["图号"], "A1")
            self.assertEqual(over[0]["数量"], 1)
            self.assertIn("当前剩余=-1", over[0]["说明"])

    def test_existing_integer_thickness_and_duplicate_history_do_not_append_again(self):
        row = self._negative_row()
        row["thickness"] = 30.0
        existing = {
            "板材号": "#OLD×1",
            "业务日期": "",
            "订单号": "ORDER-OVER",
            "图号": "A1",
            "厚度(mm)": 30,
            "数量": 1,
            "异常类型": "超加工/待核查",
            "处理结论": "正式入账，保留负数",
            "未移动原因": "",
            "处理建议": "核对订单总数量和累计加工数量；禁止自动把负数修正为0。",
            "说明": "历史首次记录",
        }
        unrelated = {
            "板材号": "#BLOCK",
            "业务日期": "",
            "订单号": "ORDER-X",
            "图号": "B1",
            "厚度(mm)": 40,
            "数量": 1,
            "异常类型": "资料异常",
            "处理结论": "阻断",
            "未移动原因": "资料不全",
            "处理建议": "核对资料",
            "说明": "保留原样",
        }

        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "累计加工台账.xlsx"
            generate_cumulative_report([row], [], [], [existing, dict(existing), unrelated], path)
            wb = load_workbook(path, data_only=True)
            ws = wb["异常记录"]
            headers = [cell.value for cell in ws[1]]
            records = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
            over = [r for r in records if r.get("异常类型") == "超加工/待核查"]
            other = [r for r in records if r.get("异常类型") == "资料异常"]

            self.assertEqual(len(over), 1)
            self.assertEqual(over[0]["厚度(mm)"], 30)
            self.assertEqual(over[0]["说明"], "历史首次记录")
            self.assertEqual(len(other), 1)
            self.assertEqual(other[0]["板材号"], "#BLOCK")


if __name__ == "__main__":
    unittest.main()
