"""未加工零件正式入口：先自愈两个正式台账，再执行原业务主流程。

自愈规则：
1. 每次都从“正在加工”根目录按正式文件名定位当前 file ID，不依赖历史固定 ID。
2. 两个正式表都存在：直接把真实 ID 注入主流程。
3. 任一正式表缺失：必须先完整、严格读取“正在加工”全部订单源；任何订单源异常都阻断重建。
4. 仅《当前待加工零件》缺失时，必须用当前订单事实 + 现有《累计加工台账》永久累计事实重建，禁止清零历史。
5. 《累计加工台账》缺失时视为一次新的当前订单基线：累计已加工从 0 开始；随后主流程继续处理
   “拆图结果”根目录现有完成文件，因此本次待入账板材仍会正常计入。
6. 正式生产模式才上传缺失表；只读测试只验证可重建性，不改 Drive。
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from modules.drive_manager import DriveManager
from modules.excel_generator import generate_cumulative_report, generate_pending_report
from modules.excel_reader import read_source_summary
from modules.fast_ledger import read_existing_ledger_fast
from modules.logger import get_logger
from modules.official_files import CUMULATIVE_NAME, PENDING_NAME, discover_official_files
from modules.process_parts import build_current_state, build_source_index, part_key

logger = get_logger()


def _test_mode() -> bool:
    return os.getenv("TEST_MODE", "false").lower() == "true"


def _verify_cumulative(path: Path) -> None:
    cumulative = load_workbook(path, read_only=True, data_only=True)
    required = {"累计加工台账", "加工流水", "板材入账记录", "异常记录"}
    if not required.issubset(set(cumulative.sheetnames)):
        raise RuntimeError(f"自愈生成的累计台账缺少工作表：{required - set(cumulative.sheetnames)}")


def _verify_pending(path: Path) -> None:
    pending = load_workbook(path, read_only=True, data_only=True)
    if "当前待加工零件" not in pending.sheetnames:
        raise RuntimeError("自愈生成的当前待加工零件缺少正式工作表")


def _read_all_sources(drive: DriveManager, temp_dir: Path) -> list[dict]:
    source_files = drive.list_order_source_files()
    logger.info(f"自愈检查：发现订单原始汇总表 {len(source_files)} 个")

    records: list[dict] = []
    errors: list[str] = []
    for idx, item in enumerate(source_files, start=1):
        container_name = item.get("order_container_name", "")
        local = temp_dir / f"heal_source_{idx}_{Path(item['name']).name}"
        drive.download_file(item["id"], str(local))
        try:
            parsed = read_source_summary(
                local,
                source_name=item["name"],
                order_container_name=container_name,
            )
        except Exception as exc:
            label = container_name or item.get("name", "未知订单")
            errors.append(f"{label}: {exc}")
            continue
        records.extend(parsed)
        logger.info(f"自愈读取订单源 {container_name or item['name']}：{len(parsed)} 条零件")

    if errors:
        preview = "；".join(errors[:8])
        if len(errors) > 8:
            preview += f"；另有 {len(errors) - 8} 个错误"
        raise RuntimeError(
            "正式台账缺失，但订单源未能全部严格读取；为防止生成漏订单台账，本次自愈已阻断。"
            f"问题：{preview}"
        )

    build_source_index(records)
    return records


def _build_zero_baseline_files(source_records: list[dict], temp_dir: Path) -> tuple[Path, Path]:
    baseline_rows = build_current_state(source_records, existing_state={})
    note = (
        "自动自愈重建：累计加工台账缺失，依据“正在加工”当前全部订单原始汇总表建立新基线；"
        "初始累计已加工为0。本次拆图结果仍会在随后正式主流程中继续核验并计入。"
    )
    cumulative_path = temp_dir / CUMULATIVE_NAME
    pending_path = temp_dir / PENDING_NAME
    generate_cumulative_report(
        baseline_rows,
        flows=[],
        board_records=[],
        anomalies=[],
        output_path=cumulative_path,
        note=note,
    )
    generate_pending_report(baseline_rows, pending_path, note=note)
    _verify_cumulative(cumulative_path)
    _verify_pending(pending_path)
    return cumulative_path, pending_path


def _build_pending_from_existing_cumulative(
    drive: DriveManager,
    cumulative_meta: dict,
    source_records: list[dict],
    temp_dir: Path,
) -> Path:
    """仅待加工表缺失时，用永久累计事实恢复动态视图，不能把累计加工量清零。"""
    local_cumulative = temp_dir / "existing_累计加工台账.xlsx"
    drive.download_file(cumulative_meta["id"], str(local_cumulative))
    existing = read_existing_ledger_fast(local_cumulative)
    source_index = build_source_index(source_records)
    rows = build_current_state(source_records, existing_state=existing["state"])

    # 当前源中已经消失但历史仍未完成的行继续冻结保留，不能在自愈过程中静默删除。
    for historical in existing.get("historical_rows", []):
        key = part_key(
            historical.get("order", ""),
            historical.get("drawing", ""),
            historical.get("thickness", 0),
            historical.get("bevel", ""),
        )
        if key not in source_index and int(historical.get("remaining", 0)) != 0:
            rows.append(historical)

    pending_path = temp_dir / PENDING_NAME
    generate_pending_report(
        rows,
        pending_path,
        note="自动自愈重建：当前待加工零件缺失，已按当前订单事实 + 现有永久累计加工事实完整恢复。",
    )
    _verify_pending(pending_path)
    return pending_path


def prepare_official_files(drive: DriveManager) -> dict[str, dict | None]:
    root_items = drive.list_children(drive.working_folder_id)
    official = discover_official_files(root_items)
    missing = [name for name, item in official.items() if item is None]
    if not missing:
        logger.info("正式台账自愈检查通过：两个正式 Excel 均存在")
        return official

    logger.warning(f"正式台账缺失，启动自愈：{missing}")
    with tempfile.TemporaryDirectory() as temp_dir_str:
        temp_dir = Path(temp_dir_str)
        source_records = _read_all_sources(drive, temp_dir)
        paths: dict[str, Path] = {}

        if CUMULATIVE_NAME in missing:
            cumulative_path, baseline_pending_path = _build_zero_baseline_files(source_records, temp_dir)
            paths[CUMULATIVE_NAME] = cumulative_path
            if PENDING_NAME in missing:
                paths[PENDING_NAME] = baseline_pending_path
        elif PENDING_NAME in missing:
            cumulative_meta = official.get(CUMULATIVE_NAME)
            if not cumulative_meta:
                raise RuntimeError("当前待加工零件缺失，但无法读取现有累计加工台账，禁止清零重建")
            paths[PENDING_NAME] = _build_pending_from_existing_cumulative(
                drive,
                cumulative_meta,
                source_records,
                temp_dir,
            )

        if _test_mode():
            logger.info(
                f"只读测试：已验证可从 {len(source_records)} 条当前订单零件安全重建 {missing}；"
                "未上传 Drive，也未执行后续正式写回。"
            )
            return official

        for name in missing:
            created = drive.upload_file(str(paths[name]), drive.working_folder_id)
            official[name] = created
            logger.info(f"正式台账自愈已新建：{name}（file_id={created['id']}）")

    refreshed = discover_official_files(drive.list_children(drive.working_folder_id))
    still_missing = [name for name, item in refreshed.items() if item is None]
    if still_missing:
        raise RuntimeError(f"正式台账自愈上传后仍未发现文件：{still_missing}")
    return refreshed


def run() -> None:
    drive = DriveManager()
    if not drive.check_config():
        raise RuntimeError("Google Drive 配置不完整")

    official = prepare_official_files(drive)
    missing_after_test = [name for name, item in official.items() if item is None]
    if missing_after_test:
        if _test_mode():
            return
        raise RuntimeError(f"正式台账自愈失败：{missing_after_test}")

    os.environ["CUMULATIVE_FILE_ID"] = str(official[CUMULATIVE_NAME]["id"])
    os.environ["PENDING_FILE_ID"] = str(official[PENDING_NAME]["id"])

    from main import run as run_business

    run_business()


if __name__ == "__main__":
    run()
