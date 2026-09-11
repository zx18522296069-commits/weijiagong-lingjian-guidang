"""正式累计加工台账与当前待加工零件 Excel 生成模块。"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from modules.process_parts import group_by_order, order_is_complete

DETAIL_HEADERS = [
    "图号", "厚度(mm)", "坡口", "长(mm)", "宽(mm)", "订单总数量",
    "零件总重量(t)", "板材号/加工来源", "累计已加工", "当前剩余",
    "待加工重量(t)", "零件状态",
]
FLOW_HEADERS = ["板材号", "内容指纹", "拆图结果文件", "图号", "厚度(mm)", "坡口", "本张板加工数量", "数据性质", "说明"]
BOARD_HEADERS = ["数据性质", "板材号", "内容指纹", "拆图结果文件", "计入件数", "状态", "归档位置"]
ANOMALY_HEADERS = ["板材号", "业务日期", "订单号", "图号", "厚度(mm)", "数量", "异常类型", "说明"]

GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
HEADER_FILL = "D9EAF7"
ORDER_FILL = "DDEBF7"
THIN = Side(style="thin", color="D9D9D9")


def _status_fill(status: str):
    if status == "已完成":
        return PatternFill("solid", fgColor=GREEN)
    if status == "部分完成":
        return PatternFill("solid", fgColor=YELLOW)
    if status == "超加工/待核查":
        return PatternFill("solid", fgColor=RED)
    return None


def _set_common_alignment(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _configure_widths(ws):
    widths = {
        "A": 29, "B": 12, "C": 10, "D": 13, "E": 13, "F": 14,
        "G": 16, "H": 25, "I": 14, "J": 14, "K": 16, "L": 16,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _write_kpis(ws, active_rows: list[dict], note: str = ""):
    grouped = group_by_order(active_rows)
    active_orders = {order: rows for order, rows in grouped.items() if not order_is_complete(rows)}
    current_weight = sum(float(r["pending_weight_t"]) for rows in active_orders.values() for r in rows)
    current_remaining = sum(int(r["remaining"]) for rows in active_orders.values() for r in rows)

    ws["A1"] = "正在进行加工订单"
    ws["E1"] = "当前未出重量（吨）"
    ws["I1"] = "当前剩余件数"
    ws["A2"] = len(active_orders)
    ws["E2"] = current_weight
    ws["I2"] = current_remaining
    for cell in (ws["A1"], ws["E1"], ws["I1"]):
        cell.font = Font(bold=True, size=12)
    ws["E2"].number_format = "0.000"

    if note:
        ws.merge_cells("A4:L4")
        ws["A4"] = note
        ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 45


def _write_order_sections(ws, rows: list[dict], *, include_completed_orders: bool):
    grouped = group_by_order(rows)
    row_cursor = 6
    for order, order_rows in grouped.items():
        completed = order_is_complete(order_rows)
        if completed and not include_completed_orders:
            continue

        pending_pieces = sum(int(r["remaining"]) for r in order_rows)
        pending_weight = sum(float(r["pending_weight_t"]) for r in order_rows)
        status = "已完成" if completed else "加工中"

        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=12)
        head = ws.cell(row_cursor, 1)
        head.value = f"订单：{order} ｜ {status} ｜ 待加工 {pending_pieces} 件 ｜ 待加工 {pending_weight:.3f} t"
        head.font = Font(bold=True, size=12)
        head.fill = PatternFill("solid", fgColor=ORDER_FILL)
        head.alignment = Alignment(horizontal="left", vertical="center")
        row_cursor += 1

        for c, header in enumerate(DETAIL_HEADERS, start=1):
            cell = ws.cell(row_cursor, c, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.border = Border(bottom=THIN)
        row_cursor += 1

        for r in order_rows:
            values = [
                r["drawing"], r["thickness"], r["bevel"], r["length"], r["width"], r["quantity"],
                r["total_weight_t"], r.get("board_sources", ""), r["processed"], r["remaining"],
                r["pending_weight_t"], r["status"],
            ]
            for c, value in enumerate(values, start=1):
                cell = ws.cell(row_cursor, c, value)
                cell.border = Border(bottom=THIN)
            ws.cell(row_cursor, 4).number_format = "0.0"
            ws.cell(row_cursor, 5).number_format = "0.0"
            ws.cell(row_cursor, 7).number_format = "0.000"
            ws.cell(row_cursor, 11).number_format = "0.000"

            # 用户确认模板：J列“当前剩余”加大、加粗、居中。
            ws.cell(row_cursor, 10).font = Font(size=14, bold=True)
            fill = _status_fill(r["status"])
            if fill:
                for col in range(1, 13):
                    ws.cell(row_cursor, col).fill = fill
            row_cursor += 1

        row_cursor += 1


def _build_main_sheet(ws, rows: list[dict], *, note: str, include_completed_orders: bool):
    _write_kpis(ws, rows, note)
    _write_order_sections(ws, rows, include_completed_orders=include_completed_orders)
    _configure_widths(ws)
    _set_common_alignment(ws)
    ws.freeze_panes = "A6"


def _write_records_sheet(ws, headers: list[str], records: list[dict]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for record in records:
        ws.append([record.get(h, "") for h in headers])
    _set_common_alignment(ws)
    for idx, header in enumerate(headers, start=1):
        width = 14
        if header in {"说明", "归档位置"}:
            width = 60 if header == "说明" else 24
        elif header in {"拆图结果文件", "图号", "订单号"}:
            width = 28
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def generate_cumulative_report(
    rows: list[dict],
    flows: list[dict],
    board_records: list[dict],
    anomalies: list[dict],
    output_path,
    *,
    note: str = "",
):
    wb = Workbook()
    ws = wb.active
    ws.title = "累计加工台账"
    _build_main_sheet(ws, rows, note=note, include_completed_orders=True)

    flow_ws = wb.create_sheet("加工流水")
    _write_records_sheet(flow_ws, FLOW_HEADERS, flows)
    board_ws = wb.create_sheet("板材入账记录")
    _write_records_sheet(board_ws, BOARD_HEADERS, board_records)
    anomaly_ws = wb.create_sheet("异常记录")
    _write_records_sheet(anomaly_ws, ANOMALY_HEADERS, anomalies)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def generate_pending_report(rows: list[dict], output_path, *, note: str = ""):
    wb = Workbook()
    ws = wb.active
    ws.title = "当前待加工零件"
    _build_main_sheet(ws, rows, note=note, include_completed_orders=False)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def generate_report(rows, output_path):
    """兼容旧接口，生成当前待加工零件。"""
    return generate_pending_report(rows, output_path)
