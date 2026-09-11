"""正式 Excel 读取与结构化解析。原始文件只读，不修改。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value, field: str, *, allow_blank: bool = False) -> float | None:
    if value in (None, ""):
        if allow_blank:
            return None
        raise ValueError(f"字段 {field} 为空")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"字段 {field} 不是有效数字: {value!r}") from exc


def _int_number(value, field: str) -> int:
    num = _number(value, field)
    rounded = int(round(num))
    if abs(num - rounded) > 1e-9:
        raise ValueError(f"字段 {field} 必须为整数: {value!r}")
    return rounded


def normalize_order_key(value: str) -> str:
    """
    将“175.26-08-05  D53K-1600A-0805”确定性归一为 D53K-1600A-0805。
    如果没有前置内部编号，则保留原值。
    """
    text = _text(value)
    if not text:
        return ""
    # 兼容人工命名时在连字符两侧误加空格，例如：
    # “THP10-8000J -0911”应视为“THP10-8000J-0911”。
    text = re.sub(r"\s*-\s*", "-", text)
    parts = [p for p in re.split(r"\s+", text) if p]
    # 订单目录允许在真实订单号后附加状态说明，例如：
    # “159.26-07-15  YT71S-2500Z-0715 已做完核算表”。
    # 只有首段是内部编号时才去掉它，并取紧随其后的真实订单号；
    # 不能取最后一段，否则会把“已做完核算表”误认为订单号。
    if len(parts) >= 2 and re.fullmatch(r"\d+\.\d{2}-\d{2}-\d{2}", parts[0]):
        return parts[1]
    return text


def normalize_drawing(value: str) -> str:
    return _text(value)


def normalize_bevel(value: str) -> str:
    return _text(value)


def _header_map(row: Iterable) -> dict[str, int]:
    result = {}
    for idx, value in enumerate(row, start=1):
        text = _text(value)
        if text:
            result[text] = idx
    return result


def _worksheet_bounds(ws) -> tuple[int, int]:
    """兼容部分第三方生成的 xlsx：read_only 模式下可能没有预写 worksheet dimension。"""
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)
    return int(ws.max_row or 0), int(ws.max_column or 0)


def _find_header(ws, required: set[str], max_rows: int = 40):
    max_row, max_column = _worksheet_bounds(ws)
    if max_row < 1 or max_column < 1:
        return None, None
    for r in range(1, min(max_row, max_rows) + 1):
        values = [ws.cell(r, c).value for c in range(1, max_column + 1)]
        texts = {_text(v) for v in values if _text(v)}
        if required.issubset(texts):
            return r, _header_map(values)
    return None, None


def _pick_header(headers: dict[str, int], *candidates: str) -> tuple[str, int] | tuple[None, None]:
    for candidate in candidates:
        if candidate in headers:
            return candidate, headers[candidate]
    return None, None


def _infer_unlabelled_bevel_column(ws, header_row: int, headers: dict[str, int],
                                   order_col: int, drawing_col: int) -> int | None:
    """识别旧模板中表头为空、但明细值明确为 P/W 的坡口列。"""
    max_row, max_column = _worksheet_bounds(ws)
    labelled_columns = set(headers.values())
    candidates: list[int] = []

    for col in range(1, max_column + 1):
        if col in labelled_columns:
            continue
        values: list[str] = []
        for row in range(header_row + 1, max_row + 1):
            order = _text(ws.cell(row, order_col).value)
            drawing = _text(ws.cell(row, drawing_col).value)
            if not order and not drawing:
                continue
            value = _text(ws.cell(row, col).value).upper()
            if value:
                values.append(value)
        if values and all(value in {"P", "W"} for value in values):
            candidates.append(col)

    return candidates[0] if len(candidates) == 1 else None


def read_source_summary(path, *, source_name: str = "", order_container_name: str = "") -> list[dict]:
    """读取订单原始汇总表，重量统一为吨；长宽、数量、坡口直接来自原始汇总表。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    records: list[dict] = []
    source_label = f"{order_container_name or '未知订单目录'} / {source_name or Path(path).name}"
    expected_order = normalize_order_key(order_container_name) if order_container_name else ""

    for ws in wb.worksheets:
        header_row, headers = _find_header(ws, {"订单号", "图号", "厚度", "件数"})
        if not header_row:
            continue

        _, order_col = _pick_header(headers, "订单号")
        _, drawing_col = _pick_header(headers, "图号")
        _, thickness_col = _pick_header(headers, "厚度", "厚度(mm)")
        _, qty_col = _pick_header(headers, "件数", "数量", "订单总数量")
        _, length_col = _pick_header(headers, "长(mm)", "长")
        _, width_col = _pick_header(headers, "宽(mm)", "宽")
        _, bevel_col = _pick_header(headers, "坡口")
        weight_name, weight_col = _pick_header(
            headers,
            "总净重(t)", "总净重（t）", "总净重(T)", "总净重（T）",
            "总净重(kg)", "总净重（kg）", "总净重(KG)", "总净重（KG）",
            "重量(吨)", "重量（吨）", "重量(t)", "重量（t）", "重量(T)", "重量（T）",
            "重量(kg)", "重量（kg）", "重量(KG)", "重量（KG）",
            "重量",
        )
        if not bevel_col and order_col and drawing_col:
            bevel_col = _infer_unlabelled_bevel_column(
                ws, header_row, headers, order_col, drawing_col
            )
        if not all([order_col, drawing_col, thickness_col, qty_col, length_col, width_col, bevel_col, weight_col]):
            raise ValueError(f"原始汇总表缺少正式必需字段: {source_label}")

        weight_is_kg = "kg" in weight_name.lower()
        max_row, _ = _worksheet_bounds(ws)

        for r in range(header_row + 1, max_row + 1):
            order_raw = _text(ws.cell(r, order_col).value)
            drawing = normalize_drawing(ws.cell(r, drawing_col).value)
            if not order_raw and not drawing:
                continue

            # 原始明细表后常附“厚度核对/合计”等汇总区。只读取与订单目录
            # 归一化订单号一致的连续明细块，防止把汇总行误当零件。
            normalized_order = normalize_order_key(order_raw) if order_raw else ""
            if expected_order and normalized_order and normalized_order != expected_order:
                if records:
                    break
                if drawing:
                    raise ValueError(
                        f"原始汇总表订单号与订单目录不一致: {source_label} 第{r}行 "
                        f"订单号={order_raw!r}"
                    )
                continue

            if not order_raw:
                continue
            if not drawing:
                raise ValueError(f"原始汇总表第{r}行图号为空: {source_label}")

            total_weight = _number(ws.cell(r, weight_col).value, weight_name or "重量")
            total_weight_t = total_weight / 1000.0 if weight_is_kg else total_weight
            records.append({
                "order_raw": order_raw,
                "order": normalized_order,
                "drawing": drawing,
                "thickness": _number(ws.cell(r, thickness_col).value, "厚度"),
                "quantity": _int_number(ws.cell(r, qty_col).value, "件数"),
                "length": _number(ws.cell(r, length_col).value, "长(mm)"),
                "width": _number(ws.cell(r, width_col).value, "宽(mm)"),
                "bevel": normalize_bevel(ws.cell(r, bevel_col).value),
                "total_weight_t": float(total_weight_t),
                "source_file": source_name or Path(path).name,
                "order_container_name": order_container_name,
                "source_row": r,
            })

        if records:
            return records

    raise ValueError(f"未找到订单汇总表正式表头: {source_label}")


def read_split_result(path, *, board_id: str, source_name: str = "") -> dict:
    """读取 `编号_完成.xlsx` 结构化明细。坡口若不存在则保留为空，由匹配阶段做唯一性检查。"""
    wb = load_workbook(path, read_only=True, data_only=True)

    for ws in wb.worksheets:
        header_row, headers = _find_header(ws, {"订单号", "图号", "厚度"})
        if not header_row:
            continue

        qty_name, split_qty_col = _pick_header(headers, "图片拆分数量", "拆分数量", "本次加工数量")
        base_qty_name, base_qty_col = _pick_header(headers, "基础件数", "基础数量")
        base_weight_name, base_weight_col = _pick_header(
            headers,
            "基础表总重量（T）", "基础表总重量(T)", "基础表总重量（t）", "基础表总重量(t)",
        )
        _, order_col = _pick_header(headers, "订单号")
        _, drawing_col = _pick_header(headers, "图号")
        _, thickness_col = _pick_header(headers, "厚度", "厚度(mm)")
        _, bevel_col = _pick_header(headers, "坡口")

        if not all([split_qty_col, base_qty_col, base_weight_col, order_col, drawing_col, thickness_col]):
            raise ValueError(f"拆图结果缺少正式必需字段: {source_name or path}")

        max_row, max_column = _worksheet_bounds(ws)
        rows: list[dict] = []
        for r in range(header_row + 1, max_row + 1):
            order_raw = _text(ws.cell(r, order_col).value)
            drawing = normalize_drawing(ws.cell(r, drawing_col).value)
            if not order_raw and not drawing:
                continue
            if not order_raw or not drawing:
                continue
            split_qty = _int_number(ws.cell(r, split_qty_col).value, qty_name)
            if split_qty <= 0:
                raise ValueError(f"拆图结果加工数量必须>0: {source_name or path} 第{r}行")
            rows.append({
                "board_id": board_id,
                "source_file": source_name or Path(path).name,
                "order_raw": order_raw,
                "order": normalize_order_key(order_raw),
                "drawing": drawing,
                "thickness": _number(ws.cell(r, thickness_col).value, "厚度"),
                "bevel": normalize_bevel(ws.cell(r, bevel_col).value) if bevel_col else "",
                "base_quantity": _int_number(ws.cell(r, base_qty_col).value, base_qty_name),
                "split_quantity": split_qty,
                "base_total_weight_t": _number(ws.cell(r, base_weight_col).value, base_weight_name),
                "source_row": r,
            })

        review = {}
        labels = {
            "图片标注重量（KG）": "image_weight_kg",
            "图片标注重量(KG)": "image_weight_kg",
            "拆分重量合计（KG）": "split_weight_kg",
            "拆分重量合计(KG)": "split_weight_kg",
            "差额绝对值（KG）": "abs_diff_kg",
            "差额绝对值(KG)": "abs_diff_kg",
            "复核结论": "review_conclusion",
        }
        for r in range(1, max_row + 1):
            for c in range(1, max_column + 1):
                label = _text(ws.cell(r, c).value)
                if label not in labels:
                    continue
                value = None
                for cc in range(c + 1, min(max_column, c + 3) + 1):
                    candidate = ws.cell(r, cc).value
                    if candidate not in (None, ""):
                        value = candidate
                        break
                if value not in (None, ""):
                    review[labels[label]] = value

        if rows:
            return {"rows": rows, "review": review}

    raise ValueError(f"未找到拆图结果正式表头或无结构化明细: {source_name or path}")


def read_existing_ledger(path) -> dict:
    """
    读取正式累计台账中的：
    - 当前累计状态（用于继续累加）
    - 主表历史零件行（用于订单源退出后永久保留已完成历史）
    - 加工流水、板材入账、异常历史
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    if "累计加工台账" not in wb.sheetnames:
        raise ValueError("累计加工台账.xlsx 缺少“累计加工台账”工作表")

    state: dict[tuple, dict] = {}
    historical_rows: list[dict] = []
    ws = wb["累计加工台账"]
    max_row, max_column = _worksheet_bounds(ws)
    if max_row < 1 or max_column < 1:
        raise ValueError("累计加工台账.xlsx 的“累计加工台账”工作表为空或无法解析范围")
    current_order = ""
    header_cols = None

    for r in range(1, max_row + 1):
        first = _text(ws.cell(r, 1).value)
        if first.startswith("订单："):
            text = first[len("订单："):]
            current_order = normalize_order_key(text.split("｜", 1)[0].strip())
            header_cols = None
            continue
        if first == "图号":
            header_cols = _header_map([ws.cell(r, c).value for c in range(1, max_column + 1)])
            continue
        if not current_order or not header_cols or not first:
            continue

        required_headers = [
            "图号", "厚度(mm)", "坡口", "长(mm)", "宽(mm)", "订单总数量",
            "零件总重量(t)", "累计已加工", "当前剩余", "待加工重量(t)", "零件状态",
        ]
        missing = [name for name in required_headers if name not in header_cols]
        if missing:
            raise ValueError(f"累计加工台账主表缺少正式字段: {missing}")

        drawing = normalize_drawing(ws.cell(r, header_cols["图号"]).value)
        if not drawing:
            continue
        thickness = float(_number(ws.cell(r, header_cols["厚度(mm)"]).value, "厚度"))
        bevel = normalize_bevel(ws.cell(r, header_cols["坡口"]).value)
        processed = _int_number(ws.cell(r, header_cols["累计已加工"]).value or 0, "累计已加工")
        board_sources = _text(ws.cell(r, header_cols.get("板材号/加工来源", 0)).value) if header_cols.get("板材号/加工来源") else ""
        remaining = _int_number(ws.cell(r, header_cols["当前剩余"]).value or 0, "当前剩余")

        key = (current_order, drawing, thickness, bevel)
        if key in state:
            raise ValueError(f"累计加工台账主表存在重复业务键: {key}")
        state[key] = {"processed": processed, "board_sources": board_sources}

        historical_rows.append({
            "order_raw": current_order,
            "order": current_order,
            "drawing": drawing,
            "thickness": thickness,
            "bevel": bevel,
            "length": _number(ws.cell(r, header_cols["长(mm)"]).value, "长(mm)"),
            "width": _number(ws.cell(r, header_cols["宽(mm)"]).value, "宽(mm)"),
            "quantity": _int_number(ws.cell(r, header_cols["订单总数量"]).value, "订单总数量"),
            "total_weight_t": float(_number(ws.cell(r, header_cols["零件总重量(t)"]).value, "零件总重量(t)")),
            "source_file": "历史累计台账",
            "order_container_name": "",
            "source_row": r,
            "board_sources": board_sources,
            "processed": processed,
            "remaining": remaining,
            "pending_weight_t": float(_number(ws.cell(r, header_cols["待加工重量(t)"]).value or 0, "待加工重量(t)")),
            "status": _text(ws.cell(r, header_cols["零件状态"]).value),
        })

    def read_sheet_records(sheet_name: str) -> list[dict]:
        if sheet_name not in wb.sheetnames:
            return []
        sheet = wb[sheet_name]
        sheet_max_row, sheet_max_column = _worksheet_bounds(sheet)
        if sheet_max_row < 1 or sheet_max_column < 1:
            return []
        headers = [_text(sheet.cell(1, c).value) for c in range(1, sheet_max_column + 1)]
        records = []
        for r in range(2, sheet_max_row + 1):
            values = [sheet.cell(r, c).value for c in range(1, sheet_max_column + 1)]
            if not any(v not in (None, "") for v in values):
                continue
            records.append({headers[i]: values[i] for i in range(len(headers)) if headers[i]})
        return records

    flows = read_sheet_records("加工流水")
    board_records = read_sheet_records("板材入账记录")
    anomalies = read_sheet_records("异常记录")
    posted_boards = {
        _text(row.get("板材号"))
        for row in board_records
        if _text(row.get("板材号"))
        and _text(row.get("状态")) not in {"作废", "未入账", "阻断"}
    }

    return {
        "state": state,
        "historical_rows": historical_rows,
        "flows": flows,
        "board_records": board_records,
        "anomalies": anomalies,
        "posted_boards": posted_boards,
    }


def read_excel(path):
    """兼容旧接口：返回第一张表的二维值，不参与正式业务逻辑。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(row) for row in ws.iter_rows(values_only=True)]
