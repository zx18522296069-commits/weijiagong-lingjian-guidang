"""累计加工台账高速只读解析。

与 excel_reader.read_existing_ledger 输出结构一致，但在 openpyxl read_only 模式下
只做一次顺序 iter_rows 扫描，避免反复 ws.cell(row, col) 导致 XML 被重复遍历。

V2 起：整单完成订单从“累计加工台账”主页面移入“已完成订单”页；读取时两页必须
合并回同一份永久累计事实，避免缓存失效或重建时丢失历史。
"""

from __future__ import annotations

import re

from openpyxl import load_workbook

from modules.excel_reader import _header_map, _int_number, _number, _text, normalize_bevel, normalize_drawing, normalize_order_key

COMPLETED_SHEET_NAME = "已完成订单"
_SCHEMA_RE = re.compile(r"weijiagong-report-schema-v(\d+)")


def _value(row: tuple, col_1based: int | None):
    if not col_1based or col_1based < 1 or col_1based > len(row):
        return None
    return row[col_1based - 1]


def _report_schema_version(wb) -> int:
    keywords = _text(getattr(wb.properties, "keywords", ""))
    match = _SCHEMA_RE.search(keywords)
    if match:
        return int(match.group(1))
    # 兼容已经有第五页、但早期未写 workbook keywords 的过渡文件。
    if COMPLETED_SHEET_NAME in wb.sheetnames:
        return 2
    return 1


def read_existing_ledger_fast(path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "累计加工台账" not in wb.sheetnames:
        raise ValueError("累计加工台账.xlsx 缺少“累计加工台账”工作表")

    state: dict[tuple, dict] = {}
    historical_rows: list[dict] = []
    required_headers = [
        "图号", "厚度(mm)", "坡口", "长(mm)", "宽(mm)", "订单总数量",
        "零件总重量(t)", "累计已加工", "当前剩余", "待加工重量(t)", "零件状态",
    ]

    def parse_detail_sheet(sheet_name: str, *, required_sheet: bool = False) -> None:
        if sheet_name not in wb.sheetnames:
            if required_sheet:
                raise ValueError(f"累计加工台账.xlsx 缺少“{sheet_name}”工作表")
            return
        ws = wb[sheet_name]
        current_order = ""
        header_cols = None
        any_row = False
        for r, values in enumerate(ws.iter_rows(values_only=True), start=1):
            any_row = True
            row = tuple(values)
            first = _text(row[0] if row else None)
            if first.startswith("订单："):
                text = first[len("订单："):]
                current_order = normalize_order_key(text.split("｜", 1)[0].strip())
                header_cols = None
                continue
            if first == "图号":
                header_cols = _header_map(row)
                missing = [name for name in required_headers if name not in header_cols]
                if missing:
                    raise ValueError(f"{sheet_name}缺少正式字段: {missing}")
                continue
            if not current_order or not header_cols or not first:
                continue

            drawing = normalize_drawing(_value(row, header_cols["图号"]))
            if not drawing:
                continue
            thickness = float(_number(_value(row, header_cols["厚度(mm)"]), "厚度"))
            bevel = normalize_bevel(_value(row, header_cols["坡口"]))
            processed = _int_number(_value(row, header_cols["累计已加工"]) or 0, "累计已加工")
            board_col = header_cols.get("板材号/加工来源")
            board_sources = _text(_value(row, board_col)) if board_col else ""
            remaining = _int_number(_value(row, header_cols["当前剩余"]) or 0, "当前剩余")

            key = (current_order, drawing, thickness, bevel)
            if key in state:
                raise ValueError(f"累计加工台账存在重复业务键: {key}（检查主表与已完成订单页是否重复）")
            state[key] = {"processed": processed, "board_sources": board_sources}
            historical_rows.append({
                "order_raw": current_order,
                "order": current_order,
                "drawing": drawing,
                "thickness": thickness,
                "bevel": bevel,
                "length": _number(_value(row, header_cols["长(mm)"]), "长(mm)"),
                "width": _number(_value(row, header_cols["宽(mm)"]), "宽(mm)"),
                "quantity": _int_number(_value(row, header_cols["订单总数量"]), "订单总数量"),
                "total_weight_t": float(_number(_value(row, header_cols["零件总重量(t)"]), "零件总重量(t)")),
                "source_file": f"历史累计台账/{sheet_name}",
                "order_container_name": "",
                "source_row": r,
                "board_sources": board_sources,
                "processed": processed,
                "remaining": remaining,
                "pending_weight_t": float(_number(_value(row, header_cols["待加工重量(t)"]) or 0, "待加工重量(t)")),
                "status": _text(_value(row, header_cols["零件状态"])),
            })

        if required_sheet and not any_row:
            raise ValueError(f"累计加工台账.xlsx 的“{sheet_name}”工作表为空或无法解析范围")

    # V1 旧文件的已完成订单还在主表；V2 新文件则分散在主表 + 已完成订单页。
    parse_detail_sheet("累计加工台账", required_sheet=True)
    parse_detail_sheet(COMPLETED_SHEET_NAME)

    def read_sheet_records(sheet_name: str) -> list[dict]:
        if sheet_name not in wb.sheetnames:
            return []
        iterator = wb[sheet_name].iter_rows(values_only=True)
        try:
            first_row = tuple(next(iterator))
        except StopIteration:
            return []
        headers = [_text(value) for value in first_row]
        records: list[dict] = []
        for values in iterator:
            row = tuple(values)
            if not any(value not in (None, "") for value in row):
                continue
            records.append({headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]})
        return records

    flows = read_sheet_records("加工流水")
    board_records = read_sheet_records("板材入账记录")
    anomalies = read_sheet_records("异常记录")
    posted_boards = {
        _text(row.get("板材号"))
        for row in board_records
        if _text(row.get("板材号"))
        and _text(row.get("状态")) not in {"作废", "未入账", "阻断"}
    }
    posted_board_keys = {
        (_text(row.get("板材号")), _text(row.get("内容指纹")))
        for row in board_records
        if _text(row.get("板材号"))
        and _text(row.get("内容指纹"))
        and _text(row.get("状态")) not in {"作废", "未入账", "阻断"}
    }
    legacy_posted_boards = {
        _text(row.get("板材号"))
        for row in board_records
        if _text(row.get("板材号"))
        and not _text(row.get("内容指纹"))
        and _text(row.get("状态")) not in {"作废", "未入账", "阻断"}
    }

    return {
        "report_schema_version": _report_schema_version(wb),
        "state": state,
        "historical_rows": historical_rows,
        "flows": flows,
        "board_records": board_records,
        "anomalies": anomalies,
        "posted_boards": posted_boards,
        "posted_board_keys": posted_board_keys,
        "legacy_posted_boards": legacy_posted_boards,
    }
